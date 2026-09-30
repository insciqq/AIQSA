"""Fixed independent checks; execute only in a bounded networkless container."""
import base64
import csv
import io
import json
import os
import re
import sys
import subprocess
import zipfile


# Candidate output is never the report channel. Run each assertion in a fresh
# bounded child, return only its numeric value over a separate pipe, and let the
# parent compare it. The outer networkless container owns the security boundary;
# this is not a general sandbox against hostile Python frame/FD introspection.
CODE_CHILD = r"""
import json
import os
import resource
import sys

resource.setrlimit(resource.RLIMIT_AS, (96 * 1024 * 1024, 96 * 1024 * 1024))
resource.setrlimit(resource.RLIMIT_CPU, (1, 1))
resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
resource.setrlimit(resource.RLIMIT_NPROC, (0, 0))
request = json.load(sys.stdin)
result_fd = int(sys.argv[1])
namespace = {"__name__": "generated"}
try:
    exec(compile(request["source"], "candidate.py", "exec"), namespace)
    function = namespace.get("sum_values")
    if not callable(function):
        raise ValueError("candidate_function_missing")
    value = function(request["values"])
    if type(value) not in (int, float):
        raise ValueError("candidate_result_invalid")
    payload = json.dumps(value, allow_nan=False).encode("ascii")
    if len(payload) > 1024:
        raise ValueError("candidate_result_invalid")
    os.write(result_fd, payload)
except BaseException:
    # SystemExit(0) from module code/function is an incomplete assertion.
    os._exit(1)
"""


def check_code(source, execution):
    if execution is not True or len(source) > 64 * 1024:
        return False
    try:
        text = source.decode("utf-8-sig")
    except UnicodeError:
        return False
    for values, expected in [([], 0), ([2, 3, 4], 9), ([-3, 7, 0], 4), ([1.5, 2.5], 4)]:
        read_fd, write_fd = os.pipe()
        try:
            os.set_blocking(read_fd, False)
            result = subprocess.run(
                [sys.executable, "-I", "-c", CODE_CHILD, str(write_fd)],
                input=json.dumps({"source": text, "values": values}),
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, text=True,
                pass_fds=(write_fd,), timeout=3, check=False,
            )
            if result.returncode != 0:
                return False
            payload = os.read(read_fd, 1025)
            if not payload or len(payload) > 1024:
                return False
            value = json.loads(payload)
            if type(value) not in (int, float) or value != expected:
                return False
        except (OSError, ValueError, subprocess.TimeoutExpired):
            return False
        finally:
            os.close(read_fd)
            os.close(write_fd)
    return True


data = json.load(sys.stdin)
case = data["case"]
if case == "fixtures":
    from PIL import Image, ImageDraw
    import openpyxl
    import docx
    import pptx
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as value:
        value.writestr("one.txt", b"alpha\n")
        value.writestr("nested/two.txt", b"beta\n")
        value.writestr("three.txt", b"gamma\n")
    screenshot = Image.new("RGB", (360, 180), "white")
    draw = ImageDraw.Draw(screenshot)
    draw.rounded_rectangle((80, 55, 280, 125), radius=12, fill="blue")
    draw.text((158, 83), "SAVE", fill="white")
    png = io.BytesIO()
    screenshot.save(png, format="PNG")
    print(json.dumps({"archive": base64.b64encode(archive.getvalue()).decode(),
                      "screenshot": base64.b64encode(png.getvalue()).decode()}))
    sys.exit(0)

files = {name: base64.b64decode(value, validate=True) for name, value in data["files"].items()}
answer = data.get("answer", "")
execution = data.get("execution", False)
office_valid = None
passed = False
try:
    if case == "squares":
        passed = execution and bool(re.search(r"\b385\b", answer))
    elif case == "total":
        passed = execution and bool(re.search(r"\b23[.,]50\b", answer))
    elif case == "inspect":
        passed = execution and bool(re.search(r"\b3\b", answer)) and bool(re.search(r"\b17\b", answer))
    elif case == "csv":
        rows = list(csv.DictReader(io.StringIO(files["records.csv"].decode("utf-8-sig"))))
        passed = rows == [{"name": "Ada", "score": "7"}, {"name": "Lin", "score": "9"}]
    elif case == "code":
        passed = check_code(files["sum_values.py"], execution)
    elif case == "zip":
        with zipfile.ZipFile(io.BytesIO(files["bundle.zip"])) as archive:
            passed = sorted(archive.namelist()) == ["alpha.txt", "beta.txt"] and archive.read("alpha.txt") == b"alpha\n" and archive.read("beta.txt") == b"beta\n"
    elif case == "chart":
        from PIL import Image
        with Image.open(io.BytesIO(files["sales.png"])) as image:
            image.verify()
        with Image.open(io.BytesIO(files["sales.png"])) as image:
            passed = image.width >= 200 and image.height >= 150 and len(image.convert("RGB").getcolors(image.width * image.height)) > 10
    elif case == "docx":
        from docx import Document
        document = Document(io.BytesIO(files["welcome.docx"]))
        office_valid = True
        text = "\n".join(p.text for p in document.paragraphs)
        passed = "Welcome" in text and "The workshop starts at 09:00. Bring a notebook." in text
    elif case == "pptx":
        from pptx import Presentation
        deck = Presentation(io.BytesIO(files["plan.pptx"]))
        office_valid = True
        texts = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
        passed = len(texts) == 3 and all(title in text and len(text.strip()) > len(title) + 5 for title, text in zip(["Plan", "Build", "Review"], texts))
    elif case == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(files["budget.xlsx"]))
        office_valid = True
        sheet = book["Budget"]
        passed = [[sheet.cell(row, column).value for column in [1, 2]] for row in range(1, 5)] == [["Item", "Amount"], ["Pens", 12], ["Paper", 8], ["Total", "=SUM(B2:B3)"]]
except Exception:
    passed = False
    if case in ["docx", "pptx", "xlsx"] and office_valid is None:
        office_valid = False
print(json.dumps({"passed": bool(passed), "officeValid": office_valid}))
